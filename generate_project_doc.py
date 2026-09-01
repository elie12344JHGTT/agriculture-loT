from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from datetime import date

OUTPUT = "agro-iot-documentation.docx"

def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=(255, 255, 255))
        shade(table.rows[0].cells[index], "1F4E5F")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table

def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)

def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)

def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.style = "No Spacing"
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(45, 45, 45)
    return paragraph

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size, color in [("Title", 28, "1F4E5F"), ("Heading 1", 18, "1F4E5F"), ("Heading 2", 13, "2F7D62"), ("Heading 3", 11, "1F4E5F")]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
p.paragraph_format.space_after = Pt(12)
r = p.add_run("AGRO IoT")
r.bold = True
r.font.size = Pt(34)
r.font.color.rgb = RGBColor(31, 78, 95)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Documentation fonctionnelle et technique")
r.font.size = Pt(18)
r.font.color.rgb = RGBColor(47, 125, 98)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
p.add_run("Plateforme de supervision agricole connectée\nFrontend React/Vite - Backend Laravel - Base de données MySQL").font.size = Pt(12)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
p.add_run(f"Document généré le {date.today().strftime('%d/%m/%Y')}").italic = True
doc.add_page_break()

# Contents
h = doc.add_heading("1. Vue d’ensemble", level=1)
doc.add_paragraph("Agro IoT est une application web de supervision agricole. Elle permet à un utilisateur autorisé de consulter l’état d’une exploitation, de suivre les mesures des capteurs, de recevoir et traiter des alertes, et de commander des actionneurs comme une pompe d’irrigation, une ventilation ou un éclairage.")
doc.add_paragraph("Le projet est séparé en deux applications qui communiquent par une API HTTP JSON :")
add_bullets(doc, [
    "Le frontend est une interface React exécutée par Vite dans le navigateur.",
    "Le backend est une API Laravel exécutée côté serveur, connectée à une base MySQL.",
    "Le dossier scripts contient une CLI PowerShell/BAT permettant de commander les actionneurs en dehors du navigateur.",
])

doc.add_heading("Objectifs fonctionnels", level=2)
add_bullets(doc, [
    "Afficher les dernières valeurs de température, humidité, CO2, luminosité et niveau d’eau.",
    "Visualiser l’évolution des mesures dans un graphique.",
    "Afficher les alertes actives et les seuils configurés.",
    "Commander les actionneurs manuellement depuis le tableau de bord ou le terminal.",
    "Consulter et exporter l’historique des mesures, alertes et actions.",
    "Gérer les utilisateurs et consulter les journaux d’accès pour un administrateur.",
])

# Architecture
doc.add_heading("2. Architecture du projet", level=1)
add_table(doc, ["Élément", "Technologie", "Responsabilité"], [
    ["agro-iot-frontend", "React + Vite + Axios", "Interface, navigation, formulaires, affichage des données et appels API."],
    ["agro-iot-backend", "Laravel + PHP", "Authentification, règles métier, API REST, autorisations et accès aux données."],
    ["Base de données", "MySQL", "Utilisateurs, profils, parcelles, capteurs, mesures, alertes, actionneurs et historiques."],
    ["scripts", "PowerShell + BAT", "Commande des actionneurs depuis un terminal Windows."],
    ["Hébergement", "Render / Vercel", "Backend Laravel sur Render et frontend React sur un hébergeur statique compatible Vite."],
], [1.5, 1.5, 4.6])
doc.add_heading("Organisation principale des dossiers", level=2)
add_table(doc, ["Dossier", "Contenu"], [
    ["frontend/src/pages", "Login, Dashboard, Historique, Alertes, Administration et Terminal."],
    ["frontend/src/components", "Composants réutilisables : cartes capteurs, panneaux, boutons, graphiques et alertes."],
    ["frontend/src/api", "Instance Axios et journalisation des actions utilisateur."],
    ["backend/routes/api.php", "Routes API d’authentification, mesures, alertes, actionneurs, historique et administration."],
    ["backend/app/Models", "Modèles Eloquent représentant les tables métier."],
    ["backend/database/migrations", "Structure des tables de la base de données."],
    ["backend/database/seeders", "Données initiales et comptes de démonstration."],
], [2.1, 5.5])

# Frontend
doc.add_heading("3. Fonctionnement du frontend", level=1)
doc.add_paragraph("Le frontend démarre avec une page d’attente, puis vérifie si une session est enregistrée dans le stockage local du navigateur. Sans session valide, il affiche la page de connexion. Après connexion, App.jsx choisit les pages accessibles selon le rôle de l’utilisateur.")
doc.add_heading("Pages disponibles", level=2)
add_table(doc, ["Page", "Fonction"], [
    ["Connexion", "Envoie l’adresse e-mail et le mot de passe au backend."],
    ["Dashboard", "Affiche les capteurs, le graphique, l’état de connexion, les alertes et les commandes rapides."],
    ["Historique", "Présente les mesures, alertes et actions avec recherche, pagination et export CSV."],
    ["Alertes", "Affiche les alertes actives, les seuils, règles automatiques et canaux de notification."],
    ["Administration", "Permet à un administrateur de gérer les utilisateurs et de consulter les accès."],
    ["Terminal", "Interprète les commandes en langage court, garde un historique local et permet ↑/↓."],
], [1.5, 6.1])
doc.add_heading("Gestion des rôles", level=2)
add_table(doc, ["Rôle", "Pages accessibles"], [
    ["Admin", "Dashboard, Historique, Alertes, Administration, Terminal"],
    ["Technicien", "Dashboard, Historique, Alertes, Terminal"],
    ["Agriculteur", "Dashboard, Historique, Alertes"],
], [1.8, 5.8])

doc.add_heading("Instance Axios", level=2)
doc.add_paragraph("Tous les appels passent par frontend/src/api/axios.js. L’URL de base est définie par VITE_API_URL ; à défaut, le frontend utilise le backend Render :")
add_code(doc, "baseURL = VITE_API_URL || https://agro-iot-backend.onrender.com")
doc.add_paragraph("Les composants ajoutent ensuite le préfixe /api, par exemple /api/measurements/latest. L’intercepteur Axios lit le token stocké dans localStorage et ajoute automatiquement l’en-tête Authorization.")

# Communication
doc.add_heading("4. Communication frontend-backend", level=1)
doc.add_paragraph("La communication suit un modèle client-serveur. Le navigateur envoie une requête HTTP au backend Laravel ; Laravel vérifie le token et les droits, lit ou modifie la base MySQL, puis renvoie une réponse JSON. Le frontend transforme cette réponse en cartes, tableaux, graphiques ou messages utilisateur.")
doc.add_heading("Schéma général", level=2)
add_code(doc, "Utilisateur\n    | action dans React\n    v\nFrontend React + Axios\n    | HTTP JSON + Authorization: Bearer <token>\n    v\nAPI Laravel\n    | validation + autorisation + logique métier\n    v\nBase MySQL\n    | données\n    ^\nRéponse JSON vers le frontend")
doc.add_heading("Principales routes utilisées", level=2)
add_table(doc, ["Méthode et route", "Utilisation"], [
    ["POST /api/auth/login", "Authentifie l’utilisateur et renvoie un token ainsi que son profil."],
    ["GET /api/auth/me", "Recharge l’utilisateur réel à partir du token après un rechargement de page."],
    ["POST /api/auth/logout", "Enregistre la déconnexion dans l’audit."],
    ["GET /api/measurements/latest", "Renvoie la dernière valeur de chaque type de capteur."],
    ["GET /api/measurements/chart", "Renvoie les labels et valeurs du graphique."],
    ["GET /api/alerts/active", "Renvoie les alertes non traitées."],
    ["PUT /api/alerts/{id}/resolve", "Marque une alerte comme traitée."],
    ["GET/PUT /api/thresholds", "Lit ou modifie les seuils critiques."],
    ["POST /api/actuators/{actuator}", "Commande irrigation, ventilation ou light."],
    ["GET /api/history/*", "Renvoie les historiques de mesures, alertes ou actions."],
    ["GET/POST/PUT/DELETE /api/users", "Gère les comptes utilisateurs pour l’administrateur."],
    ["POST /api/audit-logs", "Journalise les actions effectuées dans l’application."],
], [2.5, 5.1])

# Auth
doc.add_heading("5. Authentification et sécurité", level=1)
doc.add_heading("Connexion", level=2)
add_numbered(doc, [
    "L’utilisateur saisit son e-mail et son mot de passe dans LoginPage.jsx.",
    "Axios envoie POST /api/auth/login avec un corps JSON contenant email et password.",
    "Laravel recherche l’utilisateur, vérifie le hash du mot de passe et refuse les comptes inactifs.",
    "Laravel crée un identifiant de session d’audit et renvoie un token signé avec les informations de l’utilisateur.",
    "App.jsx conserve user, token et page active dans localStorage sous agro-iot-auth.",
    "Les requêtes suivantes reçoivent automatiquement Authorization: Bearer <token>.",
])
doc.add_paragraph("Le token est signé avec la clé de l’application, contient une date d’expiration de huit heures et identifie l’utilisateur. Le middleware AgroApiAuth refuse les requêtes protégées sans token valide. Les rôles sont contrôlés côté backend, même si le frontend masque déjà les menus non autorisés.")
doc.add_heading("Comptes de démonstration", level=2)
add_table(doc, ["Compte", "Mot de passe", "Rôle"], [
    ["admin@agri-iot.com", "password123", "Admin"],
    ["agriculteur@agri-iot.com", "password123", "Agriculteur"],
], [2.5, 2.0, 2.0])
doc.add_paragraph("Ces comptes sont créés par UserSeeder. Le Dockerfile lance ce seeder au démarrage afin de les rendre disponibles après déploiement.")

# Domain modules
doc.add_heading("6. Modules métier", level=1)
doc.add_heading("Capteurs et mesures", level=2)
doc.add_paragraph("Les capteurs sont associés à une parcelle et produisent des mesures datées. Le backend sélectionne les dernières mesures par type ; le frontend normalise les noms possibles et les affiche avec leur unité : °C, %, ppm ou lux.")
doc.add_heading("Alertes et seuils", level=2)
doc.add_paragraph("Les seuils définissent les valeurs minimales et maximales attendues. Les alertes actives sont renvoyées par Laravel. Un administrateur ou un technicien peut enregistrer un seuil ou traiter une alerte selon les règles d’autorisation.")
doc.add_heading("Actionneurs", level=2)
doc.add_paragraph("Les actionneurs représentent les équipements contrôlables : pompe d’irrigation, ventilateur et éclairage. Une commande crée une action, met à jour le statut de l’actionneur et peut être consultée dans l’historique.")
doc.add_heading("Terminal web", level=2)
doc.add_paragraph("Le terminal accepte notamment arrosage start/stop, ventilation on/off, lumiere start/stop, tout start/stop, status, date, whoami, help et history. Les commandes sont conservées dans localStorage ; les flèches haut et bas permettent de retrouver les commandes précédentes. Les résultats sont présentés avec des messages compréhensibles, sans afficher le JSON technique de l’API.")
doc.add_heading("Journal d’audit", level=2)
doc.add_paragraph("Les connexions, déconnexions, consultations, exports, commandes et erreurs importantes sont enregistrées avec l’utilisateur, le rôle, la page, l’action, le statut et la date. Cette trace permet de suivre l’utilisation du système.")

# CLI and deploy
doc.add_heading("7. CLI et déploiement", level=1)
doc.add_heading("CLI Windows", level=2)
doc.add_paragraph("Le fichier scripts/actionneurs-cli.bat sert de point d’entrée et appelle actionneurs-cli.ps1. La CLI peut être utilisée en mode menu ou directement :")
add_code(doc, ".\\scripts\\actionneurs-cli.bat arrosage start\n.\\scripts\\actionneurs-cli.bat ventilation stop\n.\\scripts\\actionneurs-cli.bat luminosite start\n.\\scripts\\actionneurs-cli.bat tout stop batch")
doc.add_paragraph("La CLI utilise AGRO_API_BASE_URL pour choisir le backend et AGRO_API_TOKEN pour fournir un token Bearer. En cas de succès, elle affiche un message utilisateur comme [OK] Arrosage activee avec succes.")
doc.add_heading("Lancement local", level=2)
add_numbered(doc, [
    "Démarrer le backend Laravel depuis agro-iot-backend avec php artisan serve, après configuration de la base de données.",
    "Définir VITE_API_URL dans le frontend si le backend local doit remplacer Render.",
    "Démarrer le frontend depuis agro-iot-frontend avec npm run dev.",
    "Ouvrir l’URL indiquée par Vite, généralement http://localhost:5173 ou http://127.0.0.1:5173.",
])
doc.add_heading("Déploiement", level=2)
doc.add_paragraph("Le backend est prévu pour Render avec Docker. Le conteneur installe les dépendances Composer, nettoie la configuration, applique les migrations, exécute UserSeeder puis démarre le serveur PHP. Les variables APP_KEY, APP_URL, DB_HOST, DB_DATABASE, DB_USERNAME, DB_PASSWORD et FRONTEND_URLS doivent être configurées dans Render. Le frontend peut être déployé séparément sur Vercel ou un hébergement statique.")

# Troubleshooting
doc.add_heading("8. Dépannage rapide", level=1)
add_table(doc, ["Symptôme", "Vérification"], [
    ["Connexion refusée", "Vérifier l’e-mail, le mot de passe, l’état actif du compte et la présence de UserSeeder sur le backend déployé."],
    ["Erreur réseau", "Vérifier VITE_API_URL, l’URL Render et que le backend répond à /api/test-connection."],
    ["Erreur CORS", "Ajouter l’origine du frontend dans FRONTEND_URLS côté Laravel."],
    ["Données absentes", "Vérifier la connexion MySQL, les migrations et les seeders de capteurs/mesures."],
    ["Commande actionneur refusée", "Vérifier le token, le rôle Admin/Technicien et la présence d’un actionneur correspondant en base."],
], [1.8, 5.8])

# Closing
doc.add_heading("Conclusion", level=1)
doc.add_paragraph("Agro IoT est une architecture web séparée et cohérente : React gère l’expérience utilisateur, Axios transporte les requêtes JSON, Laravel applique les règles métier et MySQL conserve les données. Le token Bearer relie la session frontend aux autorisations backend, tandis que l’audit et les historiques assurent la traçabilité des opérations agricoles.")

# Footer
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Agro IoT - Documentation du projet").font.size = Pt(8)

doc.core_properties.title = "Agro IoT - Documentation fonctionnelle et technique"
doc.core_properties.subject = "Fonctionnement du projet et communication frontend/backend"
doc.core_properties.author = "Agro IoT"
doc.save(OUTPUT)
print(OUTPUT)
